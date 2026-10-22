#!/usr/bin/env python3
"""Rebuild D13 from the exact editable original D12 PowerPoint.
Keep every original slide and drawing shape; report all text/geometry for review.
"""
import json
from pathlib import Path
import subprocess
from pptx import Presentation

HERE=Path(__file__).resolve().parent
SOURCE=HERE.parent/"Heating_Rate_Controller_D12_Reviewed_Styled.pptx"
OUT=HERE/"Heating_Rate_Controller_D13_Styled_DRAFT.pptx"
PDF=HERE/"Heating_Rate_Controller_D13_Styled_DRAFT.pdf"

def clean(x):
    return " ".join(x.replace("\r", " ").split())

def inspect(pres):
    slides=[]
    for sn,slide in enumerate(pres.slides,1):
        entries=[]
        for i,sh in enumerate(slide.shapes):
            text=sh.text if sh.has_text_frame else ""
            if sh.has_table:
                text=" | ".join(cell.text for row in sh.table.rows for cell in row.cells)
            entries.append({
              "index":i,"id":sh.shape_id,"name":sh.name,"kind":str(sh.shape_type),
              "x":round(sh.left/914400,3),"y":round(sh.top/914400,3),
              "w":round(sh.width/914400,3),"h":round(sh.height/914400,3),
              "text":clean(text)[:1100]
            })
        slides.append({"slide":sn,"shape_count":len(entries),"shapes":entries})
    return {"source":str(SOURCE.relative_to(HERE.parent.parent.parent)), "slide_size_inches":[round(pres.slide_width/914400,3),round(pres.slide_height/914400,3)],"slides":slides}

def main():
    assert SOURCE.exists(),SOURCE
    deck=Presentation(str(SOURCE))
    report=inspect(deck)
    (HERE/"slide_inventory.json").write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding="utf-8")
    # The first build preserves the complete original vector slides.
    deck.save(str(OUT))
    p=subprocess.run(["libreoffice","-env:UserInstallation=file:///tmp/lo_d13_2410","--headless","--convert-to","pdf","--outdir",str(HERE),str(OUT)],capture_output=True,text=True,timeout=120)
    if p.returncode!=0 or not PDF.exists():
        raise RuntimeError("PDF conversion failed: "+p.stdout+" "+p.stderr)
    assert len(Presentation(str(OUT)).slides)==len(deck.slides)==11
    print("INVENTORY READY: slides",len(deck.slides),"shapes",sum(x["shape_count"] for x in report["slides"]))
    for slide in report["slides"]:
        print("SLIDE",slide["slide"],"SHAPES",slide["shape_count"])
        for x in slide["shapes"]:
            if x["text"]:
                print("  [%02d] %s"%(x["index"],x["text"][:145]))

if __name__=="__main__":
    main()
